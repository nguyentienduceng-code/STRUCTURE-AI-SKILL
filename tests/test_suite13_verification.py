"""
test_suite13_verification.py
============================
Sprint 5 · Suite 13 — Comprehensive Structural Verification & Reporting

Kiểm thử tự động cho các tính năng nâng cấp:
  • doc_reader.py           — Eurocode materials, wind loads, scan_project_directory
  • conflict_analyzer.py   — C07 (anchorage/lap), C08 (spacing/ratio), C09 (pile capacity)
  • calc_continuous_beam.py — Continuous beam multi-span solver & TCVN 5574 design
  • calc_isolated_footing.py — Isolated footing soil pressure, punching, bending design
  • docx_report_builder.py — Professional Word (.docx) export
  • report_pipeline.py     — End-to-end multi-format reporting

Target: >= 30 checks, 0 failures.
"""

import sys
import os
import json
import unittest
import shutil
from pathlib import Path

if sys.stdout.encoding != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")

TOOLS_DIR = Path(__file__).parent.parent / "tools"
sys.path.insert(0, str(TOOLS_DIR))

import doc_reader as dr
import conflict_analyzer as ca
import calc_continuous_beam as ccb
import calc_isolated_footing as cif
import docx_report_builder as drb
import report_pipeline as rp


class TestSuite13A_DocReader(unittest.TestCase):
    """A: Nâng cấp Document Reader (Eurocode, Gió, Project Scanner)."""

    def test_A01_eurocode_concrete(self):
        text = "Sử dụng bê tông C25/30 và C30/37 theo tiêu chuẩn"
        eng, _ = dr._extract_engineering_data(text)
        grades = [m["grade"] for m in eng["materials"] if m["type"] == "concrete"]
        self.assertIn("C25/30", grades)
        c25 = next(m for m in eng["materials"] if m["grade"] == "C25/30")
        self.assertEqual(c25["fck_MPa"], 25)

    def test_A02_eurocode_structural_steel(self):
        text = "Dầm thép hình S355JR liên kết bu lông"
        eng, _ = dr._extract_engineering_data(text)
        grades = [m["grade"] for m in eng["materials"] if m["type"] == "structural_steel"]
        self.assertTrue(any("S355" in g for g in grades))
        s355 = next(m for m in eng["materials"] if "S355" in m["grade"])
        self.assertEqual(s355["fy_MPa"], 355)

    def test_A03_rebar_eurocode_astm(self):
        text = "Thép chịu kéo B500B và Grade 60"
        eng, _ = dr._extract_engineering_data(text)
        grades = [m["grade"] for m in eng["materials"] if m["type"] == "steel_rebar"]
        self.assertTrue(any("B500" in g for g in grades))

    def test_A04_wind_load_pattern(self):
        text = "Áp lực gió cơ bản W0 = 95 daN/m2, thuộc vùng gió II-A"
        eng, _ = dr._extract_engineering_data(text)
        w0_loads = [l for l in eng["loads"] if "W0" in l.get("label", "")]
        self.assertTrue(len(w0_loads) >= 1)
        self.assertEqual(w0_loads[0]["value"], 95.0)

    def test_A05_scan_project_directory(self):
        res = dr.scan_project_directory(str(TOOLS_DIR.parent))
        self.assertIn("summary", res)
        self.assertTrue(res["summary"]["total_files"] > 10)
        self.assertIn("thuyet_minh", res["inventory"])
        self.assertIn("ban_ve", res["inventory"])


class TestSuite13B_ConflictAnalyzer(unittest.TestCase):
    """B: Mở rộng Conflict Checks (C07, C08, C09)."""

    def test_B01_C07_anchorage_violation(self):
        doc = {
            "source_file": "dam.pdf",
            "raw_text_preview": "Lneo = 250 mm cho thép d25",
            "engineering_data": {
                "sections": [{"element": "D1", "b_mm": 250, "h_mm": 500}],
                "rebar_schedule": [{"dia_mm": 25.0, "count": 3}],
                "materials": [{"type": "concrete", "grade": "B25"}, {"type": "steel_rebar", "grade": "CB400V"}],
            }
        }
        conflicts, _ = ca.check_anchorage_lap(doc)
        crit = [c for c in conflicts if c["type"] == "ANCHORAGE_VIOLATION" and c["severity"] == "CRITICAL"]
        self.assertTrue(len(crit) >= 1, "Phải phát hiện neo 250mm không đủ")

    def test_B02_C07_lap_splice_violation(self):
        doc = {
            "source_file": "cot.pdf",
            "raw_text_preview": "Lnối = 350 mm",
            "engineering_data": {
                "rebar_schedule": [{"dia_mm": 25.0, "count": 4}],
                "materials": [{"type": "concrete", "grade": "B25"}, {"type": "steel_rebar", "grade": "CB400V"}],
            }
        }
        conflicts, _ = ca.check_anchorage_lap(doc)
        crit = [c for c in conflicts if c["type"] == "LAP_SPLICE_VIOLATION"]
        self.assertTrue(len(crit) >= 1, "Phải phát hiện nối chồng 350mm vi phạm TCVN")

    def test_B03_C08_rebar_spacing_congestion(self):
        doc = {
            "source_file": "dam_day.pdf",
            "raw_text_preview": "",
            "engineering_data": {
                "sections": [{"element": "Dầm D1", "b_mm": 200, "h_mm": 400}],
                "rebar_schedule": [{"dia_mm": 25.0, "count": 5}],  # 5 thanh d25 trong bề rộng 200mm -> a_clear < 25mm
                "materials": [],
            }
        }
        conflicts, _ = ca.check_rebar_spacing_and_ratio(doc)
        spacing_cf = [c for c in conflicts if c["type"] == "REBAR_SPACING_CONGESTION"]
        self.assertTrue(len(spacing_cf) >= 1, "Phải phát hiện khoảng cách cốt thép quá dày")

    def test_B04_C08_max_rebar_ratio_exceeded(self):
        doc = {
            "source_file": "dam_qua_thep.pdf",
            "raw_text_preview": "",
            "engineering_data": {
                "sections": [{"element": "Dầm D1", "b_mm": 200, "h_mm": 400}],
                "rebar_schedule": [{"dia_mm": 25.0, "count": 5}],  # As = 2454 mm2 -> mu = 3.4% > 2.5%
                "materials": [],
            }
        }
        conflicts, _ = ca.check_rebar_spacing_and_ratio(doc)
        ratio_cf = [c for c in conflicts if c["type"] == "MAX_REBAR_RATIO_EXCEEDED"]
        self.assertTrue(len(ratio_cf) >= 1, "Phải phát hiện vượt hàm lượng thép cực đại 2.5%")

    def test_B05_C09_pile_capacity_overload(self):
        doc_etabs = {"source_file": "etabs.edb", "raw_text_preview": "P_max = 850 kN"}
        doc_geo = {"source_file": "dcct.pdf", "raw_text_preview": "Rcd = 600 kN"}
        conflicts, _ = ca.check_pile_capacity_match(doc_etabs, doc_geo)
        overload = [c for c in conflicts if c["type"] == "PILE_CAPACITY_OVERLOAD"]
        self.assertTrue(len(overload) >= 1, "Phải phát hiện quá tải cọc 850kN > 600kN")


class TestSuite13C_StructuralCalculators(unittest.TestCase):
    """C: Bộ tính toán Dầm liên tục & Móng đơn BTCT."""

    def test_C01_continuous_beam_3spans(self):
        res = ccb.calc_continuous_beam(
            spans=[5.0, 6.0, 5.0],
            b=250,
            h=500,
            q_dead=15.0,
            q_live=10.0,
            concrete="B25",
            rebar="CB400-V",
        )
        self.assertEqual(res["status"], "PASSED")
        self.assertEqual(len(res["internal_forces"]["M_supports_kNm"]), 4)
        self.assertEqual(len(res["internal_forces"]["M_spans_kNm"]), 3)
        self.assertTrue(res["overall_eta"] <= 1.0)
        self.assertTrue(len(res["Calculation_Logs"]) >= 4)
        self.assertEqual(res["Standard_Method"], "TCVN 5574:2018")

    def test_C02_continuous_beam_2spans(self):
        res = ccb.calc_continuous_beam(
            spans=[6.0, 6.0],
            b=300,
            h=600,
            q_dead=20.0,
            q_live=15.0,
        )
        self.assertEqual(res["status"], "PASSED")
        # Gối giữa đối xứng mô men âm lớn: M = qL^2 / 8 = 35 * 36 / 8 = 157.5 kNm
        m_mid_sup = abs(res["internal_forces"]["M_supports_kNm"][1])
        self.assertAlmostEqual(m_mid_sup, 157.5, delta=1.0)

    def test_C03_isolated_footing_soil_pressure(self):
        res = cif.calc_isolated_footing(
            B=2.0, L=2.4, H=0.6, Df=1.5,
            N_tc=600.0, M_tc=60.0, V_tc=15.0,
            R_tc=200.0,
        )
        self.assertEqual(res["status"], "PASSED")
        self.assertTrue(res["soil_pressures"]["P_tb_kPa"] <= 200.0)
        self.assertTrue(res["soil_pressures"]["P_max_kPa"] <= 240.0)
        self.assertTrue(res["soil_pressures"]["P_min_kPa"] >= 0.0)
        self.assertTrue(res["punching_shear"]["eta_punch"] <= 1.0)
        self.assertTrue(res["reinforcement"]["As_x_req_mm2"] > 0)
        self.assertEqual(res["Standard_Method"], "TCVN 5574:2018 / TCVN 9362:2012")

    def test_C04_isolated_footing_punching(self):
        # Móng quá mỏng (H=0.3m) chịu tải lớn 1200kN -> phải cảnh báo đâm thủng
        res = cif.calc_isolated_footing(
            B=2.0, L=2.0, H=0.3, Df=1.0,
            N_tc=1200.0, M_tc=10.0, V_tc=5.0,
            R_tc=400.0,
        )
        self.assertTrue(res["punching_shear"]["eta_punch"] > 0.5)


class TestSuite13D_DocxReporting(unittest.TestCase):
    """D: Xuất Báo cáo Word (.docx) & Pipeline đầu cuối."""

    def setUp(self):
        self.test_dir = Path("./test_suite13_tmp")
        self.test_dir.mkdir(exist_ok=True)

    def tearDown(self):
        if self.test_dir.exists():
            shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_D01_docx_report_builder(self):
        dummy_data = {
            "meta": {"started_at": "2026-10-18T21:00:00", "mode": "review"},
            "read_result": {
                "source_file": "test_thuyetminh.pdf",
                "source_type": "PDF",
                "confidence_level": "HIGH",
                "engineering_data": {
                    "sections": [{"element": "Dầm D1", "b_mm": 250, "h_mm": 500}],
                    "loads": [{"label": "Tĩnh tải", "value": 15.0, "unit": "kN/m"}],
                    "materials": [{"type": "concrete", "grade": "B25"}],
                }
            },
            "conflict_result": {
                "overall_assessment": "✅ ĐẠT",
                "conflict_summary": {"total": 1, "CRITICAL": 0, "HIGH": 1, "MEDIUM": 0},
                "conflicts": [
                    {
                        "id": "C08-001", "type": "REBAR_SPACING_CONGESTION", "severity": "HIGH",
                        "param": "Khoảng cách cốt thép", "description": "a_clear < 25mm",
                        "recommendation": "Bố trí 2 lớp thép"
                    }
                ],
                "passes": [{"id": "P01", "param": "Uốn", "result": "PASS"}]
            },
            "rebar_results": [
                {"element": "Dầm D1", "rebar_notation": "4Ø25", "As_required_mm2": 1520.0, "As_provided_mm2": 1963.5, "As_ratio": 1.29, "status": "OK"}
            ],
            "calc_result": {
                "element": "Dầm D1",
                "eta_M": 0.65,
                "eta_V": 0.42,
                "overall_eta": 0.65,
                "Standard_Method": "TCVN 5574:2018",
            }
        }
        docx_out = str(self.test_dir / "report_test.docx")
        saved_path = drb.build_docx_report(dummy_data, docx_out)
        self.assertTrue(os.path.exists(saved_path))
        self.assertTrue(os.path.getsize(saved_path) > 5000)

        # Đọc lại file word kiểm tra cấu trúc
        import docx
        doc_obj = docx.Document(saved_path)
        self.assertTrue(len(doc_obj.paragraphs) >= 10)
        self.assertTrue(len(doc_obj.tables) >= 5)

    def test_D02_pipeline_e2e_docx_export(self):
        res = rp.run_pipeline(
            demo=True,
            output_dir=str(self.test_dir),
            export_docx=True,
        )
        self.assertIn("output_md", res)
        self.assertIn("output_docx", res)
        self.assertTrue(os.path.exists(res["output_md"]))
        self.assertTrue(os.path.exists(res["output_docx"]))


# ─────────────────────────────────────────────────────────────────────────────
# Test Runner
# ─────────────────────────────────────────────────────────────────────────────

def run_tests() -> bool:
    suite = unittest.TestSuite()
    for tc in [
        TestSuite13A_DocReader,
        TestSuite13B_ConflictAnalyzer,
        TestSuite13C_StructuralCalculators,
        TestSuite13D_DocxReporting,
    ]:
        suite.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(tc))

    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)
    return result.wasSuccessful()


if __name__ == "__main__":
    success = run_tests()
    sys.exit(0 if success else 1)
