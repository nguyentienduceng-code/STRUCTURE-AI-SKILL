import sys
if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')
import json
from docx import Document
from docx.shared import Pt, Inches

data = {
  "Ve_Viec": "Phê duyệt phương án thiết kế kết cấu cải tạo lấp ô sàn thông tầng tại Tầng 2 - Dự án CT PLAZA MINH CHÂU",
  "Nguoi_Nhan": "Ban Giám đốc CT Group",
  "Can_Cu_Phap_Ly": [
    "Hồ sơ bản vẽ phương án thiết kế kết cấu ngày 05/06/2026",
    "Tiêu chuẩn thiết kế kết cấu thép TCVN 5575:2024",
    "Tiêu chuẩn Tải trọng và tác động TCVN 2737:2023"
  ],
  "Hien_Trang": "Văn phòng làm việc tại tầng 2 dự án CT Plaza Minh Châu hiện có các ô thông tầng cần được lấp lại để gia tăng diện tích sử dụng. Bề mặt sàn hiện hữu là nền bê tông và gạch ceramic.",
  "Mo_Ta_Phuong_An": "Phương án thiết kế đề xuất sử dụng hệ khung thép lắp ghép (dầm chính I-300, xà gồ hộp 50x100) kết hợp với tấm lót sàn xi măng siêu nhẹ Cemboard dày 20mm.",
  "Danh_Gia_Phuong_An": [
    "Ưu điểm: Trọng lượng kết cấu cực nhẹ, giảm thiểu tải trọng tác dụng lên hệ khung hiện hữu. Phương án thi công lắp ghép khô giúp đẩy nhanh tiến độ, sạch sẽ, phù hợp không gian cải tạo.",
    "Nhược điểm: Rủi ro cháy nổ do hàn cắt tại công trường. Khả năng chống rung và cách âm kém hơn sàn bê tông thông thường."
  ],
  "Phan_Tich_So_Luoc": "Phương án lấp ô sàn sử dụng dầm thép hình I-300x150x6.5x9 kết hợp xà gồ hộp thép 50x100x1.8 và tấm sàn Cemboard dày 20mm. Bổ sung tính toán kiểm tra xà gồ hộp (nhịp 1.1m, khoảng cách 0.6m) và sức chịu uốn cục bộ của tấm Cemboard đều đạt yêu cầu hoạt tải văn phòng. Liên kết cắt giữa dầm thép và cấu kiện hiện hữu dùng bản mã 16mm khoan cấy 4 bu lông M24 (LK1). Sườn gia cường cho dầm thép được bố trí a1200 chống mất ổn định. Hệ kết cấu đáp ứng tốt khả năng chịu lực với độ võng trong giới hạn cho phép.",
  "Kien_Nghi": [
    "Thống nhất áp dụng phương án kết cấu sàn thép lấp ô thông tầng vì khối lượng nhẹ, thi công nhanh và ít ảnh hưởng tới kết cấu hiện hữu.",
    "Bắt buộc Nhà thầu đo đạc, khảo sát cao độ (cote +6.800) và kích thước lưới cột thực tế trước khi cắt thép.",
    "Bắt buộc Nhà thầu đệ trình biện pháp thi công an toàn PCCC, và dò siêu âm cốt thép trước khi khoan cấy bu lông."
  ],
  "Tac_Dong_Chi_Phi": "Tối ưu hóa nhờ sử dụng sàn thép kết hợp Cemboard, không phát sinh chi phí đúc bê tông.",
  "Tac_Dong_Tien_Do": "Thi công lắp ráp nhanh chóng, rút ngắn đáng kể thời gian so với phương án đổ sàn bê tông toàn khối.",
  "De_Nghi_Phe_Duyet": "Kính đề nghị Ban Giám đốc xem xét và phê duyệt phương án thiết kế trên để làm cơ sở triển khai thi công.",
  "Nguoi_Ky": "AI Assistant",
  "Chuc_Danh_Ky": "Kỹ sư Phòng QLTK"
}

doc = Document()

# Header
p1 = doc.add_paragraph('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM')
p1.alignment = 1
p2 = doc.add_paragraph('Độc lập - Tự do - Hạnh phúc')
p2.alignment = 1
p3 = doc.add_paragraph('-------------------')
p3.alignment = 1

# Title
title = doc.add_heading('TỜ TRÌNH', level=1)
title.alignment = 1
doc.add_paragraph(f'Về việc: {data["Ve_Viec"]}').alignment = 1
doc.add_paragraph(f'Kính gửi: {data["Nguoi_Nhan"]}').alignment = 1

# I. Căn cứ
doc.add_heading('I. CĂN CỨ PHÁP LÝ', level=2)
for cp in data["Can_Cu_Phap_Ly"]:
    doc.add_paragraph(f'- {cp}')

# II. Nội dung
doc.add_heading('II. NỘI DUNG TRÌNH DUYỆT', level=2)
doc.add_heading('1. Hiện trạng / Vấn đề', level=3)
doc.add_paragraph(data["Hien_Trang"])

doc.add_heading('2. Mô tả phương án', level=3)
doc.add_paragraph(data["Mo_Ta_Phuong_An"])

doc.add_heading('3. Phân tích kết cấu sơ lược', level=3)
doc.add_paragraph(data["Phan_Tich_So_Luoc"])
doc.add_paragraph('\n(Chi tiết xem Báo cáo phân tích kỹ thuật DeepResearch đính kèm số BXD-BCPT-MINHCHAU-REV00)')

doc.add_heading('4. Đánh giá phương án', level=3)
for dg in data["Danh_Gia_Phuong_An"]:
    doc.add_paragraph(f'- {dg}')

# III. Kiến nghị
doc.add_heading('III. KIẾN NGHỊ', level=2)
for kn in data["Kien_Nghi"]:
    doc.add_paragraph(f'- {kn}')
doc.add_paragraph(f'- Tác động chi phí: {data["Tac_Dong_Chi_Phi"]}')
doc.add_paragraph(f'- Tác động tiến độ: {data["Tac_Dong_Tien_Do"]}')

# IV. Đề nghị
doc.add_heading('IV. ĐỀ NGHỊ PHÊ DUYỆT', level=2)
doc.add_paragraph(data["De_Nghi_Phe_Duyet"])

# Signature
p_sig = doc.add_paragraph()
p_sig.alignment = 2
p_sig.add_run(f'\n\nPHÒNG QLTK\n\n\n\n{data["Nguoi_Ky"]}\n{data["Chuc_Danh_Ky"]}')

output_path = r'g:\My Drive\Google AI Studio\MINH CHAU\BXD-TTr-MINHCHAU-Phuong_An_San_Thep_V3.docx'
doc.save(output_path)
print(f"Đã tạo thành công file Word tại: {output_path}")
