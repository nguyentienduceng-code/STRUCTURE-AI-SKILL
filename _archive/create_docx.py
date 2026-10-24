import sys
if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')
import json
from docx import Document
from docx.shared import Pt, Inches

with open('data_cu_goc.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

doc = Document()

# Header
p1 = doc.add_paragraph('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM')
p1.alignment = 1
p2 = doc.add_paragraph('Độc lập - Tự do - Hạnh phúc')
p2.alignment = 1
p3 = doc.add_paragraph('-------------------')
p3.alignment = 1

# Title
title = doc.add_heading('TỜ TRÌNH', level=1)
title.alignment = 1
doc.add_paragraph(f'Về việc: {data["Ve_Viec"]}').alignment = 1
doc.add_paragraph(f'Kính gửi: {data["Nguoi_Nhan"]}').alignment = 1

# I. Căn cứ
doc.add_heading('I. CĂN CỨ PHÁP LÝ', level=2)
for cp in data["Can_Cu_Phap_Ly"]:
    doc.add_paragraph(f'- {cp}')

# II. Nội dung
doc.add_heading('II. NỘI DUNG TRÌNH DUYỆT', level=2)
doc.add_heading('1. Hiện trạng / Vấn đề', level=3)
doc.add_paragraph(data["Hien_Trang"])

doc.add_heading('2. Phân tích sơ lược', level=3)
doc.add_paragraph(data["Phan_Tich_So_Luoc"])

doc.add_heading('3. Bảng so sánh phương án', level=3)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Tiêu chí'
hdr_cells[1].text = 'PA1: Cừ đúc sẵn'
hdr_cells[2].text = 'PA2: Tự gia công'
for row in data["Bang_So_Sanh"]:
    row_cells = table.add_row().cells
    row_cells[0].text = row["Tiêu chí"]
    row_cells[1].text = row["PA1"]
    row_cells[2].text = row["PA2"]

doc.add_paragraph('\n(Chi tiết xem Báo cáo phân tích kỹ thuật đính kèm số BXD-BCPT-001-REV00)')

# III. Kiến nghị
doc.add_heading('III. KIẾN NGHỊ', level=2)
for kn in data["Kien_Nghi"]:
    doc.add_paragraph(f'- {kn}')
doc.add_paragraph(f'- Tác động chi phí: {data["Tac_Dong_Chi_Phi"]}')
doc.add_paragraph(f'- Tác động tiến độ: {data["Tac_Dong_Tien_Do"]}')

# IV. Đề nghị
doc.add_heading('IV. ĐỀ NGHỊ PHÊ DUYỆT', level=2)
doc.add_paragraph(data["De_Nghi_Phe_Duyet"])

# Signature
p_sig = doc.add_paragraph()
p_sig.alignment = 2
p_sig.add_run(f'\n\nPHÒNG QLTK\n\n\n\n{data["Nguoi_Ky"]}\n{data["Chuc_Danh_Ky"]}')

doc.save('BXD-TTr-001-Phuong_An_Cu_Goc.docx')
print("Đã tạo thành công file Word!")
