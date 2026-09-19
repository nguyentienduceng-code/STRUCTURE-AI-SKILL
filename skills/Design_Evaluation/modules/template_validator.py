import sys
if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')
import re
from docx import Document

class TemplateValidator:
    def __init__(self, template_path):
        self.doc = Document(template_path)
        self.required_tags = self._extract_tags()

    def _extract_tags(self):
        """Lấy tất cả {{tags}} từ template"""
        tags = set()
        for para in self.doc.paragraphs:
            found = re.findall(r'\{\{(\w+)\}\}', para.text)
            tags.update(found)
        # Also check tables
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        found = re.findall(r'\{\{(\w+)\}\}', para.text)
                        tags.update(found)
        return tags

    def validate_data(self, json_data):
        """Kiểm tra JSON có đủ field cho template không"""
        missing = []
        for tag in self.required_tags:
            if tag not in json_data:
                missing.append(tag)

        if missing:
            return {
                "valid": False,
                "missing_fields": missing,
                "message": f"Thiếu {len(missing)} fields trong JSON Handoff so với Template: {', '.join(missing)}"
            }

        return {"valid": True, "message": "OK - Handoff JSON hợp lệ cho template này"}
